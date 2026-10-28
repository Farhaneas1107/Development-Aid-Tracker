#!/usr/bin/env python3
"""
Extracts and enriches all project opportunity records from
Naue_Project_Opportunity_Tracker.xlsx into structured JSON and JS modules
for the Naue Project Opportunity Intelligence Portal.
"""

import json
import re
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = ROOT / "Naue_Project_Opportunity_Tracker.xlsx"
JSON_OUT = ROOT / "data" / "projects.json"
JS_OUT = ROOT / "js" / "data.js"


NAUE_PRODUCTS_CATALOG = [
    {
        "id": "Carbofol",
        "name": "Carbofol®",
        "category": "Geomembranes (HDPE / PP)",
        "color": "#0284c7",
        "shortDesc": "Impermeable polymeric geomembrane liners for basal landfill cells, capping, leachate lagoons, WWTP ponds, and irrigation canals.",
        "typicalUses": [
            "Sanitary landfill basal lining & capping",
            "Leachate & sewage sludge lagoons",
            "Farm water-harvesting ponds & desert hafirs",
            "Irrigation canal & reservoir waterproofing",
            "Secondary oil/chemical tank bund containment"
        ]
    },
    {
        "id": "Bentofix",
        "name": "Bentofix®",
        "category": "Geosynthetic Clay Liners (GCLs)",
        "color": "#0d9488",
        "shortDesc": "Needle-punched sodium bentonite mineral barrier liners providing self-healing waterproofing for landfills, dams, canals, and reservoirs.",
        "typicalUses": [
            "Secondary composite barrier beneath geomembranes",
            "Dumpsite closure & landfill capping barriers",
            "Dam embankments, barrages & canal seepage control",
            "Water storage ponds & stormwater attenuation basins"
        ]
    },
    {
        "id": "Secutex",
        "name": "Secutex®",
        "category": "Nonwoven Geotextiles",
        "color": "#2563eb",
        "shortDesc": "Mechanically bonded nonwoven geotextiles for filtration ('filter cloth' behind gabions), cushion protection of geomembranes, separation, and scour control.",
        "typicalUses": [
            "Filter cloth behind gabion walls & weir aprons",
            "Heavy-duty cushion protection for Carbofol® liners",
            "Road subbase & utility trench separation/filtration",
            "River/wadi bridge abutment & culvert scour protection",
            "Dewatering / sediment containment tubes"
        ]
    },
    {
        "id": "Secugrid",
        "name": "Secugrid®",
        "category": "Geogrids (Soil Reinforcement)",
        "color": "#4f46e5",
        "shortDesc": "Monolithic flat-bar PET/PP geogrids for high-modulus road base stabilization, bridge approach MSE walls, railway trackbeds, and steep slopes.",
        "typicalUses": [
            "Highway, rural & municipal road base reinforcement",
            "Bridge approach embankments & MSE retaining walls",
            "Heavy-duty landfill access roads & transfer ramps",
            "Solar/wind farm access roads & substation yards",
            "Landfill capping steep slope veneer reinforcement"
        ]
    },
    {
        "id": "Combigrid",
        "name": "Combigrid®",
        "category": "Composite Geogrid + Geotextile",
        "color": "#7c3aed",
        "shortDesc": "Geocomposite combining a monolithic Secugrid® reinforcement geogrid with an integrated Secutex® nonwoven geotextile for soft subgrades.",
        "typicalUses": [
            "Combined reinforcement, separation & filtration in one pass",
            "Weak/marshy road subgrades & urban street rehabilitation",
            "Heavy-axle waste truck access roads"
        ]
    },
    {
        "id": "Secudrain",
        "name": "Secudrain®",
        "category": "Drainage Geocomposites",
        "color": "#0891b2",
        "shortDesc": "Three-dimensional drainage core laminated with Secutex® filter geotextiles for high-capacity liquid and gas transmission under load.",
        "typicalUses": [
            "Landfill leachate collection & leak detection layers",
            "Landfill gas (LFG) venting beneath capping systems",
            "Rainwater & structural subsurface drainage networks"
        ]
    },
    {
        "id": "Secumat",
        "name": "Secumat®",
        "category": "Erosion Control Mats",
        "color": "#16a34a",
        "shortDesc": "Three-dimensional permanent erosion control and turf reinforcement mats protecting soil slopes and hydraulic channels against surface runoff.",
        "typicalUses": [
            "Reconstructed 1.5:1 earth irrigation canal side slopes",
            "Landfill capping vegetated surface erosion control",
            "Flood embankments, dam slopes & coastal banks"
        ]
    },
    {
        "id": "Secutex Soft Rock",
        "name": "Secutex® Soft Rock",
        "category": "Coastal & Hydraulic Sand Containers",
        "color": "#ea580c",
        "shortDesc": "High-durability nonwoven geotextile sand containers engineered for coastal protection, revetments, groynes, and marine scour defense.",
        "typicalUses": [
            "Climate-resilient coastal protection revetments",
            "Marine & port scour protection structures",
            "Dune stabilization & shoreline erosion defense"
        ]
    }
]


# Curated enrichment metadata keyed by Excel row number (5..35)
CURATED_META = {
    5: {
        "shortName": "Wadi Tuban Flood Response & Canal Rehab (ERA)",
        "countries": ["Yemen"],
        "sectors": ["Irrigation, Canals & Dams"],
        "donors": ["UNDP", "KfW"],
        "salesChannel": "Direct Civil Works Supply",
        "timelineTag": "Active Tender (Drawings Mar 2026)",
        "budget": "8 Canal Interventions (>3,397 ha)",
        "stats": {"total": 8, "high": 4, "medium": 4, "low": 0, "actionable": 8},
        "mapLocations": [
            {"name": "LT-C1: Obar Al-Hadharim Canal", "lat": 13.083992, "lng": 44.887034, "endLat": 13.076197, "endLng": 44.910712, "precision": "exact", "note": "High Opp: 1.5:1 earth canal side slope reconstruction"},
            {"name": "LT-C2: Al-Tha'alab Weir & Gates", "lat": 13.101580, "lng": 44.864202, "endLat": 13.081361, "endLng": 44.899504, "precision": "exact", "note": "Med-High Opp: Weir rehab & 25 hydraulic gates scour filter"},
            {"name": "LT-C3: Obar Al-Ihssan Canal", "lat": 13.144920, "lng": 44.833240, "endLat": 13.145310, "endLng": 44.834551, "precision": "exact", "note": "Med-High Opp: Guide walls & foundation filtration"},
            {"name": "LT-C4: Obar Fareed Intake & Gabions", "lat": 13.148700, "lng": 44.822570, "endLat": 13.148449, "endLng": 44.822818, "precision": "exact", "note": "Very High Opp: Explicit FILTER CLOTH behind stepped gabion walls & embankments"},
            {"name": "LT-C5: Obar Al-Makharij Gabions", "lat": 13.143140, "lng": 44.830860, "endLat": 13.143061, "endLng": 44.830835, "precision": "exact", "note": "Very High Opp: Stepped gabion walls & aprons w/ filter geotextile"},
            {"name": "Int. 6: Falaj Al-Ninoh", "lat": 13.122183, "lng": 44.847393, "endLat": 13.122023, "endLng": 44.848076, "precision": "exact", "note": "Med-High Opp: Hydraulic structures & scour protection"},
            {"name": "Int. 7: Al-Fiyah & Al-Faqih", "lat": 13.053400, "lng": 44.861720, "endLat": 13.042630, "endLng": 44.864340, "precision": "exact", "note": "High Opp: 100m+ earth canal 1.5:1 slope reconstruction"},
            {"name": "Int. 8: Al-Fardhah", "lat": 13.017546, "lng": 44.964927, "endLat": 13.018560, "endLng": 44.965545, "precision": "exact", "note": "Med-High Opp: Gate station & concrete guide wall filtration"}
        ]
    },
    6: {
        "shortName": "Zahrat Al-Finjan & Al-Minya Sanitary Landfills",
        "countries": ["Palestine (West Bank & Gaza)"],
        "sectors": ["Sanitary Landfills & Waste"],
        "donors": ["World Bank"],
        "salesChannel": "Consultant / BoQ Specification",
        "timelineTag": "SPNs Jul 2025 & May 2026",
        "budget": "3 Packages (2 Major Landfills)",
        "stats": {"total": 3, "high": 2, "medium": 0, "low": 1, "actionable": 2},
        "mapLocations": [
            {"name": "Zahrat Al-Finjan Sanitary Landfill (Jenin)", "lat": 32.388, "lng": 35.197, "precision": "approximate", "note": "New cells, old cell rehab, LFG & leachate systems (Design + Supervision)"},
            {"name": "Al-Minya Sanitary Landfill (Bethlehem/Hebron)", "lat": 31.621, "lng": 35.228, "precision": "approximate", "note": "New cells, rehab, LFG & leachate treatment design"}
        ]
    },
    7: {
        "shortName": "Iraq Social Fund for Development (SFD Community Works)",
        "countries": ["Iraq"],
        "sectors": ["Roads, Highways & Bridges", "Water, WWTP & Stormwater"],
        "donors": ["World Bank"],
        "salesChannel": "Direct Civil Works Supply",
        "timelineTag": "Rolling Procurement through Jun 2026",
        "budget": ">US$ 42.5M (US$ 29.7M Civil/Water)",
        "stats": {"total": 86, "high": 20, "medium": 12, "low": 54, "actionable": 32},
        "mapLocations": [
            {"name": "Iraq SFD Nationwide (Baghdad PMO)", "lat": 33.3152, "lng": 44.3661, "precision": "regional", "note": "86 sub-projects across 18 governorates (>150 villages)"},
            {"name": "Basra & Maysan Rural Roads/Water", "lat": 30.5085, "lng": 47.7804, "precision": "regional", "note": "Soft/marshy subgrade road reinforcement & compact water units"},
            {"name": "Sulaymaniyah, Erbil & Duhok Mountain Roads", "lat": 35.5558, "lng": 45.4351, "precision": "regional", "note": "Mountainous village road networks & rainwater drainage"},
            {"name": "Babylon & Al-Najaf Roads & Bridges", "lat": 32.4682, "lng": 44.4196, "precision": "regional", "note": "BAB-03 roads & NJF-02-2 vehicular/pedestrian bridges"}
        ]
    },
    8: {
        "shortName": "EODP-AF: MOWR Anbar Water Works & Road Rehab Phase 1",
        "countries": ["Iraq"],
        "sectors": ["Irrigation, Canals & Dams", "Roads, Highways & Bridges"],
        "donors": ["World Bank"],
        "salesChannel": "Consultant / BoQ Specification",
        "timelineTag": "Active through Apr 2028",
        "budget": "US$ 4.94M+ (Intl QCBS Supervision)",
        "stats": {"total": 3, "high": 2, "medium": 1, "low": 0, "actionable": 3},
        "mapLocations": [
            {"name": "Anbar MOWR Hydraulic & Canal Corridor", "lat": 33.4259, "lng": 43.2992, "precision": "regional", "note": "Euphrates barrage, irrigation canal rehab & Road Rehab Phase 1"}
        ]
    },
    9: {
        "shortName": "MOA Agriculture Resilience & Water Harvesting (ARDI)",
        "countries": ["Jordan"],
        "sectors": ["Irrigation, Canals & Dams"],
        "donors": ["World Bank"],
        "salesChannel": "Consultant / BoQ Specification",
        "timelineTag": "Active Phases through Jun 2027",
        "budget": "3 WB Trust Funds (27 Activities)",
        "stats": {"total": 27, "high": 0, "medium": 6, "low": 21, "actionable": 6},
        "mapLocations": [
            {"name": "Hamad Basin Desert Water Harvesting", "lat": 32.0, "lng": 38.5, "precision": "approximate", "note": "Desert rangeland hafirs & water harvesting basin study"},
            {"name": "Zarqa Basin Farm Water Harvesting", "lat": 32.1, "lng": 36.0, "precision": "approximate", "note": "Farm-level rainwater harvesting ponds & MOA civil experts"}
        ]
    },
    10: {
        "shortName": "WAJ Non-Revenue Water (NRW), Energy Efficiency & RE",
        "countries": ["Jordan"],
        "sectors": ["Water, WWTP & Stormwater", "Energy & Industrial"],
        "donors": ["World Bank"],
        "salesChannel": "Consultant / BoQ Specification",
        "timelineTag": "Active Schedule through May 2029",
        "budget": ">US$ 14.3M (Consultancies)",
        "stats": {"total": 6, "high": 0, "medium": 3, "low": 3, "actionable": 3},
        "mapLocations": [
            {"name": "Water Authority of Jordan (National WAJ Program)", "lat": 31.9632, "lng": 35.9080, "precision": "regional", "note": "NRW storage, WWTP energy upgrades & RE solar site stabilization"}
        ]
    },
    11: {
        "shortName": "Blueing the Black Sea GEF Regional Project (BBSEA)",
        "countries": ["Türkiye", "Georgia", "Moldova", "Ukraine"],
        "sectors": ["Coastal & Marine", "Sanitary Landfills & Waste"],
        "donors": ["World Bank", "GEF"],
        "salesChannel": "Programmatic / Downstream",
        "timelineTag": "Approved Jan 2023 · 1-Yr Pilot Windows",
        "budget": "GEF Regional Grant + 4 Windows",
        "stats": {"total": 5, "high": 0, "medium": 4, "low": 1, "actionable": 4},
        "mapLocations": [
            {"name": "Black Sea Basin Regional Hub", "lat": 43.0, "lng": 34.0, "precision": "approximate", "note": "Eco-innovation pilot grants in Türkiye, Georgia, Moldova & pre-feasibility study"}
        ]
    },
    12: {
        "shortName": "Gaza SWM: Sofa Leachate Facility & Deir Al Balah Closure",
        "countries": ["Palestine (West Bank & Gaza)"],
        "sectors": ["Sanitary Landfills & Waste"],
        "donors": ["World Bank"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "Doc. 2017–2022 · Execution Pending",
        "budget": "7 Procurement Packages (GSWMP & AF)",
        "stats": {"total": 7, "high": 4, "medium": 0, "low": 3, "actionable": 4},
        "mapLocations": [
            {"name": "Sofa Sanitary Landfill & Leachate Facility", "lat": 31.2431, "lng": 34.3472, "precision": "approximate", "note": "Leachate pre-treatment lagoons & earthworks/containment cells"},
            {"name": "Deir Al Balah Sanitary Landfill / Dumpsite", "lat": 31.4042, "lng": 34.3792, "precision": "approximate", "note": "Full dumpsite closure & capping civil works"}
        ]
    },
    13: {
        "shortName": "GSWMP Sofa Landfill Access Road & Transfer Stations",
        "countries": ["Palestine (West Bank & Gaza)"],
        "sectors": ["Roads, Highways & Bridges", "Sanitary Landfills & Waste"],
        "donors": ["World Bank"],
        "salesChannel": "Direct Civil Works Supply",
        "timelineTag": "Doc. 2017–2021 · Execution Pending",
        "budget": "6 Civil & Supervision Packages",
        "stats": {"total": 6, "high": 2, "medium": 0, "low": 4, "actionable": 2},
        "mapLocations": [
            {"name": "Khan Younis Transfer Station & Sofa Road", "lat": 31.3444, "lng": 34.3028, "precision": "approximate", "note": "Heavy-duty landfill access road & elevated tipping ramp"},
            {"name": "Rafah Solid Waste Transfer Station", "lat": 31.2886, "lng": 34.2519, "precision": "approximate", "note": "Transfer station civil works & MSE ramp potential"}
        ]
    },
    14: {
        "shortName": "EODP National Highways, Culverts & River Bridges",
        "countries": ["Iraq"],
        "sectors": ["Roads, Highways & Bridges"],
        "donors": ["World Bank"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "Doc. 2016–2022 · Pending Packages",
        "budget": "~US$ 78M+ Documented Works",
        "stats": {"total": 18, "high": 15, "medium": 3, "low": 0, "actionable": 15},
        "mapLocations": [
            {"name": "Mosul 2nd, 4th & 5th Tigris Bridges", "lat": 36.3500, "lng": 43.1167, "precision": "approximate", "note": "$19.75M river bridge packages & approach embankments"},
            {"name": "Samarra Diversion Bridge & Al-Dhuluiya", "lat": 34.1833, "lng": 43.8667, "precision": "approximate", "note": "$15.1M Samarra & $4.79M Al-Dhuluiya bridge approaches"},
            {"name": "Diyala Highways & Narin/Imam Ways Bridges", "lat": 33.7500, "lng": 44.6333, "precision": "approximate", "note": "Al-Khalis–Adheem ($10.8M) & Naft Khanna–Mandli ($9.23M) roads"}
        ]
    },
    15: {
        "shortName": "EODP Municipal WWTP, Stormwater & Water Purification",
        "countries": ["Iraq"],
        "sectors": ["Water, WWTP & Stormwater"],
        "donors": ["World Bank"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "Doc. 2017–2020 · Pending Works",
        "budget": "15 Municipal Water/WWTP Packages",
        "stats": {"total": 15, "high": 1, "medium": 4, "low": 10, "actionable": 5},
        "mapLocations": [
            {"name": "Tikrit: Al-Dhebaee WWTP (C2C)", "lat": 34.6000, "lng": 43.6833, "precision": "approximate", "note": "High Opp: Detailed design & lagoon/sludge bed lining"},
            {"name": "Jalula & As-Sadiya Stormwater Networks (W07)", "lat": 34.2667, "lng": 45.1500, "precision": "approximate", "note": "Stormwater networks, pump station & WTP rehab"},
            {"name": "Mosul Right Bank & Al-Gwear Water Works", "lat": 36.3333, "lng": 43.1333, "precision": "approximate", "note": "NIN-W02 ($649k) & NIN-W03 ($243k) water purification"}
        ]
    },
    16: {
        "shortName": "EODP Al-Anbar Municipal Streets & Market Zones Rehab",
        "countries": ["Iraq"],
        "sectors": ["Roads, Highways & Bridges"],
        "donors": ["World Bank"],
        "salesChannel": "Direct Civil Works Supply",
        "timelineTag": "Doc. 2019 · Pending Execution",
        "budget": "5 Municipal RFB Civil Packages",
        "stats": {"total": 5, "high": 0, "medium": 5, "low": 0, "actionable": 5},
        "mapLocations": [
            {"name": "Al-Ramadi Streets & Central Market (ANB-W01/W02)", "lat": 33.4258, "lng": 43.2992, "precision": "approximate", "note": "Urban street subbase & market paving"},
            {"name": "Al-Falluja & Khaldiyah Market Streets (ANB-W03/W07)", "lat": 33.3500, "lng": 43.7833, "precision": "approximate", "note": "Commercial street rehabilitation"},
            {"name": "Heet Old Commercial Market (ANB-W08)", "lat": 33.6333, "lng": 42.8167, "precision": "approximate", "note": "Market zone pavement stabilization"}
        ]
    },
    17: {
        "shortName": "Jordan Third Municipal Services & Social Resilience AF",
        "countries": ["Jordan"],
        "sectors": ["Roads, Highways & Bridges", "Sanitary Landfills & Waste", "Water, WWTP & Stormwater"],
        "donors": ["World Bank"],
        "salesChannel": "Programmatic / Downstream",
        "timelineTag": "Approved Apr 2020 · Municipal Grants",
        "budget": "Multi-Municipal IPF Grant Program",
        "stats": {"total": 10, "high": 0, "medium": 3, "low": 7, "actionable": 3},
        "mapLocations": [
            {"name": "Jordan Host Municipalities (Amman MOPIC/MOLA HQ)", "lat": 31.9539, "lng": 35.9106, "precision": "approximate", "note": "Municipal roads, solid waste cells & stormwater sub-tenders"}
        ]
    },
    18: {
        "shortName": "Water, Sanitation, Urban & Energy Priority Investment TC",
        "countries": ["Multi-Country / Regional"],
        "sectors": ["Water, WWTP & Stormwater", "Roads, Highways & Bridges"],
        "donors": ["EBRD / Multilateral"],
        "salesChannel": "Consultant / BoQ Specification",
        "timelineTag": "Pre-Signing Feasibility Stage",
        "budget": "EUR 149,000 Pre-Signing TC",
        "stats": {"total": 2, "high": 0, "medium": 2, "low": 0, "actionable": 2},
        "mapLocations": []
    },
    19: {
        "shortName": "Tajikistan Feed the Future Agriculture & Water / TAWON",
        "countries": ["Tajikistan"],
        "sectors": ["Irrigation, Canals & Dams"],
        "donors": ["USAID"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "USAID Pipeline · Pending Rollout",
        "budget": "US$ 20M–50M (2 Water Contracts)",
        "stats": {"total": 4, "high": 0, "medium": 2, "low": 2, "actionable": 2},
        "mapLocations": [
            {"name": "Tajikistan Feed the Future Districts (Dushanbe HQ)", "lat": 38.5597, "lng": 68.7870, "precision": "approximate", "note": "2 parallel $10M–$25M irrigation canal & water reservoir contracts"}
        ]
    },
    20: {
        "shortName": "USAID Dedicated A&E & Infrastructure Contracts (4 Missions)",
        "countries": ["Egypt", "Lebanon", "Liberia", "Thailand"],
        "sectors": ["Roads, Highways & Bridges", "Irrigation, Canals & Dams"],
        "donors": ["USAID"],
        "salesChannel": "Consultant / BoQ Specification",
        "timelineTag": "USAID Multi-Year A&E Pipeline",
        "budget": "US$ 11.65M–30.5M (4 Contracts)",
        "stats": {"total": 4, "high": 0, "medium": 4, "low": 0, "actionable": 4},
        "mapLocations": [
            {"name": "Egypt Infrastructure A&E (5 Governorates)", "lat": 30.0444, "lng": 31.2357, "precision": "regional", "note": "$1M–$3.99M local A&E monitoring & specification"},
            {"name": "Lebanon USAID Mission A&E Contract", "lat": 33.8938, "lng": 35.5018, "precision": "regional", "note": "$10M–$24.99M nationwide A&E support"}
        ]
    },
    21: {
        "shortName": "Yemen Sana'a & Aden Sanitation Mapping & Crisis Recovery",
        "countries": ["Yemen"],
        "sectors": ["Water, WWTP & Stormwater"],
        "donors": ["World Bank"],
        "salesChannel": "Programmatic / Downstream",
        "timelineTag": "Diagnostic Precursor ('Towards 2030')",
        "budget": "US$ 100M ECRP + IDA18 Diagnostics",
        "stats": {"total": 2, "high": 0, "medium": 0, "low": 2, "actionable": 1},
        "mapLocations": [
            {"name": "Sana'a Sanitation Chain Mapping", "lat": 15.3694, "lng": 44.1910, "precision": "approximate", "note": "Diagnostic precursor for future WWTP lagoon rehab"},
            {"name": "Aden Sanitation Chain Mapping", "lat": 12.7855, "lng": 45.0181, "precision": "approximate", "note": "Diagnostic precursor for future wastewater ponds"}
        ]
    },
    22: {
        "shortName": "EODP Heet 132kV Substation Civil Works & Grid Equipment",
        "countries": ["Iraq"],
        "sectors": ["Energy & Industrial"],
        "donors": ["World Bank"],
        "salesChannel": "Direct Civil Works Supply",
        "timelineTag": "Doc. 2016–2021 · 1 Civil Lot",
        "budget": "~60 Sub-Contracts ($14.89M+ Goods)",
        "stats": {"total": 60, "high": 0, "medium": 0, "low": 60, "actionable": 1},
        "mapLocations": [
            {"name": "Heet 132kV Substation Civil Works (Anbar)", "lat": 33.6333, "lng": 42.8167, "precision": "approximate", "note": "Transformer oil bund containment & substation yard stabilization"}
        ]
    },
    23: {
        "shortName": "USAID Urban WASH, Education & Climate Adaptation Pipeline",
        "countries": ["Multi-Country / Regional"],
        "sectors": ["Coastal & Marine", "Water, WWTP & Stormwater"],
        "donors": ["USAID"],
        "salesChannel": "Programmatic / Downstream",
        "timelineTag": "USAID Multi-Year Forecast",
        "budget": "US$ 50M+ Multi-Mission Awards",
        "stats": {"total": 8, "high": 0, "medium": 0, "low": 8, "actionable": 2},
        "mapLocations": [
            {"name": "Barbados Coastal Climate Adaptation", "lat": 13.1939, "lng": -59.5432, "precision": "regional", "note": "$10M–$24.99M climate adaptation (potential coastal erosion control)"},
            {"name": "Nigeria Urban WASH Support", "lat": 9.0765, "lng": 7.3986, "precision": "regional", "note": "$1M–$3.99M urban water utility investment planning"}
        ]
    },
    24: {
        "shortName": "MENA Desert Ecosystems Knowledge Sharing (MENA-DELP)",
        "countries": ["Multi-Country / Regional"],
        "sectors": ["Policy / Non-Civil"],
        "donors": ["World Bank", "GEF"],
        "salesChannel": "Non-Civil (No Action)",
        "timelineTag": "Category C Technical Assistance",
        "budget": "US$ 1.00M GEF TA Grant",
        "stats": {"total": 5, "high": 0, "medium": 0, "low": 5, "actionable": 0},
        "mapLocations": []
    },
    25: {
        "shortName": "USAID Non-Infrastructure Portfolio (~55 Social/Health Awards)",
        "countries": ["Multi-Country / Regional"],
        "sectors": ["Policy / Non-Civil"],
        "donors": ["USAID"],
        "salesChannel": "Non-Civil (No Action)",
        "timelineTag": "Non-Construction Social Portfolio",
        "budget": "~55 Non-Civil Awards (21 Countries)",
        "stats": {"total": 55, "high": 0, "medium": 0, "low": 55, "actionable": 0},
        "mapLocations": []
    },
    26: {
        "shortName": "WB&G Multi-Donor Solid Waste, North Gaza & Hebron Sewage",
        "countries": ["Palestine (West Bank & Gaza)"],
        "sectors": ["Sanitary Landfills & Waste", "Water, WWTP & Stormwater", "Roads, Highways & Bridges"],
        "donors": ["World Bank", "USAID", "EU / Bilateral (SIDA, AFD)"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "Multi-Phase Portfolio · Pending Works",
        "budget": "9 Major Multi-Donor Programs",
        "stats": {"total": 9, "high": 6, "medium": 2, "low": 1, "actionable": 8},
        "mapLocations": [
            {"name": "North Gaza Emergency Sewage & Percolation Basins", "lat": 31.5381, "lng": 34.5042, "precision": "regional", "note": "Emergency WWTP & infiltration/percolation basins"},
            {"name": "Hebron Regional Wastewater & Southern WB SWM", "lat": 31.5326, "lng": 35.0998, "precision": "regional", "note": "Sanitary landfills & regional wastewater treatment"}
        ]
    },
    27: {
        "shortName": "Jordan National Solid Waste Management Strategy (EU NSWMS)",
        "countries": ["Jordan"],
        "sectors": ["Sanitary Landfills & Waste"],
        "donors": ["EU (ENI)"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "ENI 21-Contract Framework",
        "budget": "21 Direct Management Contracts",
        "stats": {"total": 21, "high": 8, "medium": 4, "low": 9, "actionable": 12},
        "mapLocations": [
            {"name": "Northern & Central Jordan Landfills (Irbid, Mafraq, Zarqa)", "lat": 32.5568, "lng": 35.8469, "precision": "regional", "note": "Sanitary landfill cells, leachate ponds & dumpsite capping across 8 governorates"}
        ],
        "inferredProducts": ["Carbofol", "Bentofix", "Secudrain", "Secutex"]
    },
    28: {
        "shortName": "Iran Gharesoo Zaringol Irrigation & Fenderesk Dam (IsDB)",
        "countries": ["Iran"],
        "sectors": ["Irrigation, Canals & Dams"],
        "donors": ["Islamic Development Bank (IsDB)"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "IsDB General Procurement Notice (NCB)",
        "budget": "Dam, Flood System & Irrigation Network",
        "stats": {"total": 6, "high": 3, "medium": 1, "low": 2, "actionable": 4},
        "mapLocations": [
            {"name": "Fenderesk Dam & Gharesoo Zaringol (Golestan)", "lat": 36.9800, "lng": 54.9200, "precision": "regional", "note": "Dam construction, flood prevention embankments & irrigation/drainage network"}
        ],
        "inferredProducts": ["Carbofol", "Bentofix", "Secutex", "Secumat", "Secugrid"]
    },
    29: {
        "shortName": "Iraq WB Corridors, Irrigation ($140M), Rail ($200M) & Dams",
        "countries": ["Iraq"],
        "sectors": ["Roads, Highways & Bridges", "Irrigation, Canals & Dams", "Water, WWTP & Stormwater"],
        "donors": ["World Bank"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "Major National Infrastructure Pipeline",
        "budget": "US$ 440M+ Documented Pipeline",
        "stats": {"total": 7, "high": 5, "medium": 2, "low": 0, "actionable": 7},
        "mapLocations": [
            {"name": "Dokan & Derbandikhan Dams / KRG Corridors", "lat": 35.9542, "lng": 44.9528, "precision": "regional", "note": "Hydropower dam rehab & transport corridors"},
            {"name": "Southern Marshlands & Babil Irrigation ($140M)", "lat": 31.8500, "lng": 46.8000, "precision": "regional", "note": "Irrigation/drainage networks, 100 km roads & $200M railway trackbed"}
        ],
        "inferredProducts": ["Carbofol", "Bentofix"]
    },
    30: {
        "shortName": "Lebanon Water Augmentation, Lake Qaraoun & Beirut Transport",
        "countries": ["Lebanon"],
        "sectors": ["Irrigation, Canals & Dams", "Sanitary Landfills & Waste", "Roads, Highways & Bridges"],
        "donors": ["World Bank"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "CDR / CTB Multi-Project Pipeline",
        "budget": "US$ 375M+ (Qaraoun $55M, Transport $200M)",
        "stats": {"total": 8, "high": 3, "medium": 3, "low": 2, "actionable": 6},
        "mapLocations": [
            {"name": "Lake Qaraoun Pollution Prevention (Bekaa Valley)", "lat": 33.5681, "lng": 35.6947, "precision": "regional", "note": "$55M wastewater & solid waste/leachate containment"},
            {"name": "Greater Beirut Water Augmentation & Urban Transport", "lat": 33.8886, "lng": 35.4955, "precision": "regional", "note": "Water supply reservoirs/dams & $200M urban transport corridors"}
        ],
        "inferredProducts": ["Carbofol", "Bentofix", "Secugrid", "Secutex"]
    },
    31: {
        "shortName": "Iraq JICA Waste Landfills, Ports, Bridges & Oil Containment",
        "countries": ["Iraq"],
        "sectors": ["Sanitary Landfills & Waste", "Coastal & Marine", "Roads, Highways & Bridges", "Energy & Industrial"],
        "donors": ["JICA"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "JICA ODA Loan & Design Schemes",
        "budget": "7 Major JICA ODA Sector Programs",
        "stats": {"total": 7, "high": 5, "medium": 2, "low": 0, "actionable": 7},
        "mapLocations": [
            {"name": "Basrah Refinery, Ports & Sanitary Landfill ODA", "lat": 30.5150, "lng": 47.8100, "precision": "regional", "note": "Port rehab, refinery secondary containment & leachate-stopping sanitary landfills"},
            {"name": "Samawah Bridges & Roads Construction", "lat": 31.3189, "lng": 45.2806, "precision": "regional", "note": "JICA Loan Aid bridge & highway construction"}
        ],
        "inferredProducts": ["Carbofol", "Bentofix", "Secugrid", "Secutex", "Secutex Soft Rock"]
    },
    32: {
        "shortName": "Jordan WB CPF: Badia Ponds, Wastewater Reuse & Solar/Wind",
        "countries": ["Jordan"],
        "sectors": ["Irrigation, Canals & Dams", "Water, WWTP & Stormwater", "Energy & Industrial"],
        "donors": ["World Bank"],
        "salesChannel": "Programmatic / Downstream",
        "timelineTag": "CPF Portfolio & Physical Sub-Projects",
        "budget": "7 Portfolio Programs (WB / IFC)",
        "stats": {"total": 7, "high": 0, "medium": 4, "low": 3, "actionable": 4},
        "mapLocations": [
            {"name": "Jordan Badia Ecosystem & Tafila Wind/Solar", "lat": 30.8337, "lng": 35.6161, "precision": "regional", "note": "Badia water-harvesting ponds, solar/wind access roads & Red Sea–Dead Sea"}
        ],
        "inferredProducts": ["Carbofol", "Bentofix", "Secugrid", "Secutex"]
    },
    33: {
        "shortName": "Oman FAO Integrated Water Resources Management Pilot",
        "countries": ["Oman"],
        "sectors": ["Irrigation, Canals & Dams"],
        "donors": ["FAO"],
        "salesChannel": "Programmatic / Downstream",
        "timelineTag": "FAO Pilot Allocation",
        "budget": "US$ 300,000 IWRM Pilot",
        "stats": {"total": 4, "high": 0, "medium": 1, "low": 3, "actionable": 1},
        "mapLocations": [
            {"name": "Oman Water-Stressed Farm Pilot Regions", "lat": 23.5880, "lng": 58.3829, "precision": "regional", "note": "Farm-level irrigation pond & IWRM pilot"}
        ],
        "inferredProducts": ["Carbofol", "Secutex"]
    },
    34: {
        "shortName": "Yemen Climate Resilience of Coastal Communities (CRCC)",
        "countries": ["Yemen"],
        "sectors": ["Coastal & Marine"],
        "donors": ["World Bank"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "US$ 20M Grant · Pending Rollout",
        "budget": "US$ 20.0M CSCF Grant",
        "stats": {"total": 3, "high": 1, "medium": 1, "low": 1, "actionable": 2},
        "mapLocations": [
            {"name": "Hodeida Coastal Protection Infrastructure", "lat": 14.7979, "lng": 42.9530, "precision": "regional", "note": "Small/medium-scale coastal protection revetments & marine filters"},
            {"name": "Aden Coastal Resilience Sites", "lat": 12.7855, "lng": 45.0181, "precision": "regional", "note": "Coastal erosion defense — Secutex® Soft Rock sand containers"}
        ]
    },
    35: {
        "shortName": "Jordan KfW Water-Wastewater, Sludge & Agricultural Reuse",
        "countries": ["Jordan"],
        "sectors": ["Water, WWTP & Stormwater", "Irrigation, Canals & Dams"],
        "donors": ["KfW"],
        "salesChannel": "Direct + Specification",
        "timelineTag": "€55M Program · International Bidding",
        "budget": "€55.0M (€50M KfW Loan)",
        "stats": {"total": 4, "high": 2, "medium": 2, "low": 0, "actionable": 4},
        "mapLocations": [
            {"name": "North Amman (Shafa Badran) & As Salt Sewerage/Sludge", "lat": 32.0636, "lng": 35.8856, "precision": "regional", "note": "Sewage sludge drying beds, digestate lagoons & sewer networks"},
            {"name": "Zarqa (Sukhneh) & Irbid Agricultural Reuse Reservoirs", "lat": 32.1333, "lng": 36.0667, "precision": "regional", "note": "Large-scale treated wastewater storage reservoirs for agriculture"}
        ]
    }
}


def normalize_relevance(rel_text: str, sheet2_rating: str = None):
    t = (sheet2_rating or rel_text or "").strip()
    upper = rel_text.upper() if rel_text else ""

    if sheet2_rating:
        s = sheet2_rating.strip()
        mapping = {
            "Very High": ("Very High", 6),
            "High": ("High", 5),
            "Medium-to-High": ("Medium-to-High", 4),
            "Medium": ("Moderate", 3),
            "Low-to-Medium": ("Low-to-Medium", 2),
            "Low": ("Low", 1),
        }
        if s in mapping:
            return mapping[s][0], mapping[s][1], s

    if upper.startswith("VERY HIGH"):
        return "Very High", 6, "Very High"
    if upper.startswith("HIGH TO MODERATE") or upper.startswith("MEDIUM TO HIGH") or upper.startswith("MEDIUM-TO-HIGH"):
        return "Medium-to-High", 4, "Medium-to-High"
    if upper.startswith("HIGH"):
        return "High", 5, "High"
    if upper.startswith("MODERATE") or upper.startswith("MEDIUM"):
        return "Moderate", 3, "Moderate"
    if upper.startswith("LOW TO MEDIUM") or upper.startswith("LOW-TO-MEDIUM"):
        return "Low-to-Medium", 2, "Low-to-Medium"
    if upper.startswith("LOW"):
        return "Low", 1, "Low"
    if upper.startswith("NONE"):
        return "None", 0, "None"
    return "Moderate", 3, t[:24]


def extract_products(scope_text: str, rel_text: str, sheet2_products: str, inferred=None):
    combined = f"{scope_text or ''}\n{rel_text or ''}\n{sheet2_products or ''}"
    found = []
    explicit_flags = {}

    # Check Secutex Soft Rock first
    if re.search(r"Secutex[®\s]*Soft\s*Rock", combined, re.I):
        found.append("Secutex Soft Rock")
        explicit_flags["Secutex Soft Rock"] = True

    patterns = [
        ("Carbofol", r"Carbofol"),
        ("Bentofix", r"Bentofix"),
        ("Secutex", r"Secutex(?!\s*Soft\s*Rock)"),
        ("Secugrid", r"Secugrid"),
        ("Combigrid", r"Combigrid"),
        ("Secudrain", r"Secudrain"),
        ("Secumat", r"Secumat"),
    ]
    for pid, pat in patterns:
        if re.search(pat, combined, re.I):
            found.append(pid)
            explicit_flags[pid] = True

    if inferred:
        for pid in inferred:
            if pid not in found:
                found.append(pid)
                explicit_flags[pid] = False

    return found, explicit_flags


def parse_scope(scope_text: str):
    if not scope_text:
        return "", []
    text = scope_text.strip()
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]

    lead = ""
    items = []
    opp_callout = ""

    for ln in lines:
        if ln.upper().startswith("OPPORTUNITIES FOR NAUE"):
            opp_callout = re.sub(r"^OPPORTUNITIES FOR NAUE\s*", "", ln, flags=re.I).lstrip(": ")
            continue
        if ln.upper().startswith("WHAT IS HAPPENING:"):
            cleaned = re.sub(r"^WHAT IS HAPPENING:\s*", "", ln, flags=re.I)
            # Check if inline (1), (2), (3) exist
            parts = re.split(r"(?=\(\d+\)\s+|\([ivx]+\)\s+)", cleaned)
            if len(parts) > 1:
                raw_lead = parts[0].strip().rstrip(":;")
                raw_lead = re.sub(r"\s*(Activities|Projects)\s+happening\s+include\s*$", "", raw_lead, flags=re.I).strip()
                lead = raw_lead if raw_lead else "Key program activities & procurement packages:"
                for p in parts[1:]:
                    m = re.match(r"^\(([^)]+)\)\s*(.*)", p.strip())
                    if m:
                        body_txt = re.sub(r"^and\s+", "", m.group(2).strip().rstrip(";"), flags=re.I)
                        items.append({"num": m.group(1), "title": "", "body": body_txt[0].upper() + body_txt[1:] if body_txt else ""})
                    else:
                        items.append({"num": "•", "title": "", "body": p.strip()})
            else:
                lead = re.sub(r"\s*(Activities|Projects)\s+happening\s+include:?\s*$", "", cleaned, flags=re.I).strip()
            continue

        # Numbered item: "1. Title: Body"
        m_num = re.match(r"^(\d+)\.\s*([^:]+):\s*(.+)$", ln)
        if m_num:
            items.append({
                "num": m_num.group(1),
                "title": m_num.group(2).strip(),
                "body": m_num.group(3).strip()
            })
            continue

        # Bullet item: "• Title: Body"
        m_bul = re.match(r"^[•\-]\s*([^:]+):\s*(.+)$", ln)
        if m_bul:
            items.append({
                "num": "•",
                "title": m_bul.group(1).strip(),
                "body": m_bul.group(2).strip()
            })
            continue

        if not lead:
            # Check if single paragraph with semicolon-separated clauses (like Dev Aid_3)
            if ";" in ln and len(ln) > 180:
                clauses = [c.strip() for c in ln.split(";") if c.strip()]
                lead = clauses[0] + "."
                for idx, cl in enumerate(clauses, 1):
                    cl_clean = re.sub(r"^and\s+", "", cl, flags=re.I)
                    items.append({"num": str(idx), "title": "", "body": cl_clean[0].upper() + cl_clean[1:]})
            else:
                lead = ln.rstrip(":")
        else:
            items.append({"num": "•", "title": "", "body": ln.lstrip("•- ").strip()})

    return lead, items, opp_callout


def parse_timeline(tl_text: str):
    if not tl_text:
        return "", []
    lines = [ln.strip() for ln in tl_text.strip().split("\n") if ln.strip()]
    headline = lines[0].rstrip(":") if lines else ""
    milestones = []
    if len(lines) > 1:
        for ln in lines[1:]:
            milestones.append(ln.lstrip("•- ").strip())
    elif ";" in headline:
        parts = [p.strip() for p in headline.split(";") if p.strip()]
        headline = parts[0]
        milestones = parts[1:]
    return headline, milestones


def parse_relevance_details(rel_text: str):
    if not rel_text:
        return "", "", []
    lines = [ln.strip() for ln in rel_text.strip().split("\n") if ln.strip()]
    headline = lines[0] if lines else ""
    why_matters = ""
    breakdown = []

    for ln in lines[1:]:
        clean = ln.lstrip("•- ").strip()
        if clean.lower().startswith("why it matters:"):
            why_matters = clean[len("why it matters:"):].strip()
        elif clean.lower().startswith("breakdown of opportunities"):
            continue
        else:
            # Classify sub-opportunity bullet tier
            tier = "info"
            low_c = clean.lower()
            if any(k in low_c for k in ["very high", "high opportunity", "high opportunities", "high relevance", "direct supply"]):
                tier = "high"
            elif any(k in low_c for k in ["no-relevance", "no relevance", "low / no", "none"]):
                tier = "low"
            elif any(k in low_c for k in ["medium", "moderate", "niche", "downstream"]):
                tier = "medium"
            breakdown.append({"tier": tier, "text": clean})

    if not why_matters and len(lines) == 1 and ":" in headline:
        # e.g. "High: Solid waste management projects..."
        parts = headline.split(":", 1)
        headline = parts[0].strip() + " Relevance"
        why_matters = parts[1].strip()

    return headline, why_matters, breakdown


def extract_tender_codes(ref_text: str):
    if not ref_text:
        return []
    # Extract recognizable IDs/codes
    codes = []
    patterns = [
        r"Package\s*\d+",
        r"SPN-[A-Za-z]+\s*\d{4}",
        r"IBRD\s*/\s*\d+",
        r"TF\s*/\s*[A-Z0-9]+",
        r"P\d{6}",
        r"EODP-[A-Z0-9/\-\.]+",
        r"GZ-MDLF-[A-Z0-9\-]+",
        r"SOL-\d+-\d+-\d+",
        r"ENI\s*/\s*\d{4}\s*/\s*[\d\-]+",
        r"CRIS\s*[A-Za-z0-9/\- ]+",
    ]
    for pat in patterns:
        for m in re.findall(pat, ref_text):
            val = re.sub(r"\s+", " ", m.strip())
            if val not in codes and len(codes) < 8:
                codes.append(val)
    if not codes:
        first_line = ref_text.split("\n")[0].strip()
        if len(first_line) <= 48:
            codes.append(first_line)
    return codes


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    s1 = wb["Project Opportunity Tracker"]
    s2 = wb["Opportunity Dashboard & KPIs"]

    # Map Sheet 2 KPI rows (rows 5..11 correspond to Sheet 1 rows 5..11)
    sheet2_by_row = {}
    for idx, r in enumerate(range(5, 12), start=5):
        vals = [s2.cell(r, c).value for c in range(1, s2.max_column + 1)]
        sheet2_by_row[idx] = {
            "refShort": vals[0],
            "countryShort": vals[1],
            "shortName": vals[2],
            "agencyShort": vals[3],
            "rating": vals[4],
            "totalSub": vals[5],
            "highOpps": vals[6],
            "medOpps": vals[7],
            "lowOpps": vals[8],
            "actionableOpps": vals[9],
            "productFit": vals[10],
        }

    projects = []
    folder_labels = {
        "Development Aid": "Batch 1 · Development Aid",
        "Development Aid_2": "Batch 2 · Development Aid 2",
        "Development Aid_3": "Batch 3 · Development Aid 3",
    }

    for r in range(5, s1.max_row + 1):
        vals = [s1.cell(r, c).value for c in range(1, s1.max_column + 1)]
        if not any(vals):
            continue

        folder = (vals[0] or "").strip()
        country_raw = (vals[1] or "").strip()
        location_raw = (vals[2] or "").strip()
        coords_raw = (vals[3] or "").strip()
        ref_raw = (vals[4] or "").strip()
        title_raw = (vals[5] or "").strip()
        org_raw = (vals[6] or "").strip()
        scope_raw = (vals[7] or "").strip()
        timeline_raw = (vals[8] or "").strip()
        rel_raw = (vals[9] or "").strip()

        s2_info = sheet2_by_row.get(r, {})
        meta = CURATED_META.get(r, {})

        rel_tier, rel_rank, rel_badge = normalize_relevance(rel_raw, s2_info.get("rating"))
        products, explicit_product_map = extract_products(
            scope_raw, rel_raw, s2_info.get("productFit"), meta.get("inferredProducts")
        )
        scope_lead, scope_items, opp_callout = parse_scope(scope_raw)
        tl_headline, tl_milestones = parse_timeline(timeline_raw)
        rel_headline, why_matters, rel_breakdown = parse_relevance_details(rel_raw)
        tender_codes = extract_tender_codes(ref_raw)

        seq = len(projects) + 1
        proj = {
            "id": f"NAUE-{seq:02d}",
            "seq": seq,
            "excelRow": r,
            "folder": folder,
            "folderLabel": folder_labels.get(folder, folder),
            "shortName": s2_info.get("shortName") or meta.get("shortName") or title_raw[:65],
            "title": title_raw,
            "countryRaw": country_raw,
            "countries": meta.get("countries", [country_raw]),
            "locationRaw": location_raw,
            "coordinatesRaw": coords_raw,
            "mapLocations": meta.get("mapLocations", []),
            "tenderRefRaw": ref_raw,
            "tenderCodes": tender_codes,
            "procuringOrgRaw": org_raw,
            "agencyShort": s2_info.get("agencyShort") or ", ".join(meta.get("donors", [])),
            "donors": meta.get("donors", []),
            "sectors": meta.get("sectors", ["Water, WWTP & Stormwater"]),
            "salesChannel": meta.get("salesChannel", "Direct + Specification"),
            "technicalScopeRaw": scope_raw,
            "scopeLead": scope_lead,
            "scopeItems": scope_items,
            "scopeOppCallout": opp_callout,
            "timelineRaw": timeline_raw,
            "timelineTag": meta.get("timelineTag") or tl_headline[:48],
            "timelineHeadline": tl_headline,
            "timelineMilestones": tl_milestones,
            "relevanceRaw": rel_raw,
            "relevanceTier": rel_tier,
            "relevanceRank": rel_rank,
            "relevanceHeadline": rel_headline,
            "whyItMatters": why_matters,
            "relevanceBreakdown": rel_breakdown,
            "products": products,
            "explicitProductMap": explicit_product_map,
            "primaryProductSummary": s2_info.get("productFit") or "",
            "budget": meta.get("budget", "Not specified"),
            "stats": meta.get("stats", {"total": 1, "high": 0, "medium": 1, "low": 0, "actionable": 1}),
        }
        projects.append(proj)

    payload = {
        "documentTitle": "NAUE GEOSYNTHETICS — INTERNATIONAL DEVELOPMENT AID PROJECT OPPORTUNITY TRACKER",
        "documentSubtitle": "Unified Interactive Intelligence Portal across All 31 Development Aid Programs (Batches 1–3) | Technical Scope, Geosynthetic Product Fit, Site Locations & Team Review Workflow",
        "sourceFile": "Naue_Project_Opportunity_Tracker.xlsx",
        "generatedAt": "2026-10-28",
        "executionContextNote": "All 31 documented programs—including multi-year tenders with earlier scheduled dates that have not yet started physical execution—are active in this tracker for review, consultant specification, and civil contractor pursuit.",
        "productsCatalog": NAUE_PRODUCTS_CATALOG,
        "projects": projects,
    }

    JSON_OUT.parent.mkdir(parents=True, exist_ok=True)
    JS_OUT.parent.mkdir(parents=True, exist_ok=True)

    JSON_OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    JS_OUT.write_text(
        "// Auto-generated from Naue_Project_Opportunity_Tracker.xlsx via scripts/build_data.py\n"
        "window.NAUE_INITIAL_DATA = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )

    print(f"Successfully extracted {len(projects)} projects.")
    total_sub = sum(p["stats"]["total"] for p in projects)
    total_act = sum(p["stats"]["actionable"] for p in projects)
    print(f"Total sub-projects/tenders: {total_sub} | Total actionable Naue opportunities: {total_act}")
    tiers = {}
    for p in projects:
        tiers[p["relevanceTier"]] = tiers.get(p["relevanceTier"], 0) + 1
    print("Relevance tiers:", tiers)


if __name__ == "__main__":
    main()
