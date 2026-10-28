#!/usr/bin/env python3
"""
Local Web & API Server for the Naue Geosynthetics Project Opportunity Tracker.
Serves the interactive portal and provides persistent disk storage for team
reviews plus live Excel (.xlsx) import and reviewed-Excel export.
"""

import argparse
import io
import json
import mimetypes
import subprocess
import sys
import threading
import webbrowser
from datetime import datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DATA_JSON = ROOT / "data" / "projects.json"
REVIEWS_JSON = ROOT / "data" / "team_reviews.json"
EXCEL_FILE = ROOT / "Naue_Project_Opportunity_Tracker.xlsx"
BUILD_SCRIPT = ROOT / "scripts" / "build_data.py"


def load_reviews():
    if REVIEWS_JSON.exists():
        try:
            return json.loads(REVIEWS_JSON.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def merge_project_review(existing: dict, incoming: dict) -> dict:
    if not existing:
        return incoming
    if not incoming:
        return existing

    by_id = {}
    for idx, item in enumerate(existing.get("entries") or existing.get("comments") or []):
        eid = item.get("id") or f"{item.get('author', 'rev')}_{item.get('isoTime') or item.get('date') or idx}"
        by_id[eid] = item
    for idx, item in enumerate(incoming.get("entries") or incoming.get("comments") or []):
        nid = item.get("id") or f"{item.get('author', 'rev')}_{item.get('isoTime') or item.get('date') or idx}"
        by_id[nid] = item

    merged_entries = sorted(
        by_id.values(),
        key=lambda x: str(x.get("isoTime") or ""),
        reverse=True,
    )
    newer = incoming if str(incoming.get("updatedAt") or "") >= str(existing.get("updatedAt") or "") else existing
    merged_checklist = {**(existing.get("checklist") or {}), **(incoming.get("checklist") or {})}

    return {
        "status": newer.get("status") or existing.get("status") or "unreviewed",
        "starred": bool(newer.get("starred")),
        "assignee": newer.get("assignee") or existing.get("assignee") or "",
        "checklist": merged_checklist,
        "updatedAt": newer.get("updatedAt") or datetime.now().isoformat(),
        "updatedBy": newer.get("updatedBy") or existing.get("updatedBy") or "",
        "entries": merged_entries,
    }


def save_reviews(incoming_map: dict):
    current = load_reviews()
    for pid, rec in incoming_map.items():
        if isinstance(rec, dict):
            current[pid] = merge_project_review(current.get(pid), rec)
    REVIEWS_JSON.parent.mkdir(parents=True, exist_ok=True)
    REVIEWS_JSON.write_text(json.dumps(current, ensure_ascii=False, indent=2), encoding="utf-8")
    return current


def generate_reviewed_excel_bytes() -> bytes:
    """Create an enriched .xlsx file combining project data + team review state."""
    projects_payload = json.loads(DATA_JSON.read_text(encoding="utf-8"))
    projects = projects_payload.get("projects", [])
    reviews = load_reviews()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reviewed Opportunity Tracker"

    # Title banner
    ws.merge_cells("A1:P1")
    ws["A1"] = "NAUE GEOSYNTHETICS — REVIEWED INTERNATIONAL DEVELOPMENT AID OPPORTUNITY TRACKER"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="0B2545")
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:P2")
    ws["A2"] = f"Exported {datetime.now().strftime('%Y-%m-%d %H:%M')} | All {len(projects)} Programs Across Batches 1–3 with Team Review Status, Product Fit & Notes"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="1E3A5F")
    ws["A2"].fill = PatternFill("solid", fgColor="E0F2FE")
    ws["A2"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[2].height = 20

    headers = [
        "Project ID",
        "Review Status",
        "Starred",
        "Assigned Lead",
        "Relevance Tier",
        "Actionable Opps",
        "Total Sub-Tenders",
        "Sales Channel",
        "Applicable Naue Products",
        "Country / Region",
        "Short Program Name",
        "Procuring / Funding Agency",
        "Documented Budget",
        "Timeline / Stage",
        "Team Review Notes",
        "Source Batch",
    ]

    header_fill = PatternFill("solid", fgColor="00549F")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    thin_side = Side(style="thin", color="CBD5E1")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    ws.row_dimensions[4].height = 24
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = border

    tier_fills = {
        "Very High": "D1FAE5",
        "High": "DBEAFE",
        "Medium-to-High": "CFFAFE",
        "Moderate": "FEF3C7",
        "Low-to-Medium": "E0E7FF",
        "Low": "F1F5F9",
        "None": "F8FAFC",
    }

    status_labels = {
        "unreviewed": "Unreviewed",
        "pursuing": "Pursuing (Priority)",
        "consultant_spec": "Consultant Spec",
        "contractor_bid": "Target Contractors",
        "monitoring": "Monitoring / Hold",
        "not_applicable": "Not Applicable",
    }

    all_individual_rows = []
    for idx, p in enumerate(projects, start=5):
        rev = reviews.get(p["id"], {})
        status_code = rev.get("status", "unreviewed")
        entries_list = rev.get("entries") or rev.get("comments") or []
        notes_str = "\n".join(
            f"[{c.get('displayTime') or c.get('date') or ''}] {c.get('author', 'Reviewer')} ({c.get('verdictLabel') or status_labels.get(status_code, '')}): {c.get('text', '')}"
            for c in entries_list
        )
        for c in entries_list:
            all_individual_rows.append([
                c.get("displayTime") or c.get("date") or c.get("isoTime") or "",
                p["id"],
                p["shortName"],
                p["countryRaw"],
                c.get("author", "Reviewer"),
                c.get("verdictLabel") or status_labels.get(c.get("verdict") or status_code, ""),
                c.get("text", ""),
            ])
        if rev.get("quickNote"):
            notes_str = (rev["quickNote"] + ("\n" + notes_str if notes_str else "")).strip()

        row_vals = [
            p["id"],
            status_labels.get(status_code, status_code),
            "★ Yes" if rev.get("starred") else "",
            rev.get("assignee", ""),
            p["relevanceTier"],
            p["stats"]["actionable"],
            p["stats"]["total"],
            p["salesChannel"],
            ", ".join(prod + "®" for prod in p["products"]) if p["products"] else "None",
            p["countryRaw"],
            p["shortName"],
            p["procuringOrgRaw"],
            p["budget"],
            p["timelineTag"],
            notes_str,
            p["folder"],
        ]

        ws.row_dimensions[idx].height = 42
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = Font(name="Calibri", size=10, bold=(c_idx in (1, 5)))
            cell.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
            cell.border = border
            if c_idx == 5 and p["relevanceTier"] in tier_fills:
                cell.fill = PatternFill("solid", fgColor=tier_fills[p["relevanceTier"]])

    widths = [12, 20, 10, 18, 16, 15, 15, 24, 34, 24, 38, 36, 22, 26, 42, 18]
    for c_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = w

    ws.freeze_panes = "A5"

    # Sheet 2: All Individual Team Reviews Log (Multi-Reviewer)
    ws_rev = wb.create_sheet("All Individual Reviews")
    rev_headers = [
        "Timestamp (Auto)",
        "Project ID",
        "Program Short Name",
        "Country",
        "Reviewer Name",
        "Reviewer Stage / Verdict",
        "Review Assessment / Notes",
    ]
    for col_idx, h in enumerate(rev_headers, start=1):
        c = ws_rev.cell(row=1, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.border = border
    for r_idx, r_vals in enumerate(all_individual_rows, start=2):
        for c_idx, val in enumerate(r_vals, start=1):
            c = ws_rev.cell(row=r_idx, column=c_idx, value=val)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border
    for c_idx, w in enumerate([22, 12, 38, 22, 20, 24, 65], start=1):
        ws_rev.column_dimensions[get_column_letter(c_idx)].width = w
    ws_rev.freeze_panes = "A2"

    # Also copy full verbatim details sheet
    ws2 = wb.create_sheet("Full Verbatim Details")
    full_headers = [
        "Project ID",
        "Folder Name",
        "Project Country",
        "Project Location",
        "Project Coordinates",
        "Tender Reference",
        "Tender Title",
        "Procuring Organisation",
        "Technical Scope",
        "Procurement and Implementation Timeline",
        "Relevance to Naue",
    ]
    for col_idx, h in enumerate(full_headers, start=1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    for idx, p in enumerate(projects, start=2):
        vals = [
            p["id"],
            p["folder"],
            p["countryRaw"],
            p["locationRaw"],
            p["coordinatesRaw"],
            p["tenderRefRaw"],
            p["title"],
            p["procuringOrgRaw"],
            p["technicalScopeRaw"],
            p["timelineRaw"],
            p["relevanceRaw"],
        ]
        for c_idx, v in enumerate(vals, start=1):
            c = ws2.cell(row=idx, column=c_idx, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border
    for c_idx, w in enumerate([12, 18, 22, 32, 32, 28, 40, 34, 55, 36, 50], start=1):
        ws2.column_dimensions[get_column_letter(c_idx)].width = w
    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TrackerRequestHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def log_message(self, fmt, *args):
        sys.stdout.write(f"[Naue Portal] {self.address_string()} - {fmt % args}\n")

    def _send_json(self, payload: dict, status: int = 200):
        raw = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(raw)

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/health":
            return self._send_json({"ok": True, "server": "naue-opportunity-portal"})

        if parsed.path == "/api/state":
            projects_data = {}
            if DATA_JSON.exists():
                projects_data = json.loads(DATA_JSON.read_text(encoding="utf-8"))
            return self._send_json({
                "ok": True,
                "data": projects_data,
                "reviews": load_reviews(),
            })

        if parsed.path == "/api/export-excel":
            try:
                xlsx_bytes = generate_reviewed_excel_bytes()
                fname = f"Naue_Reviewed_Opportunity_Tracker_{datetime.now().strftime('%Y%m%d')}.xlsx"
                self.send_response(200)
                self.send_header(
                    "Content-Type",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                self.send_header("Content-Disposition", f'attachment; filename="{fname}"')
                self.send_header("Content-Length", str(len(xlsx_bytes)))
                self.end_headers()
                self.wfile.write(xlsx_bytes)
                return
            except Exception as exc:
                return self._send_json({"ok": False, "error": str(exc)}, status=500)

        return super().do_GET()

    def do_POST(self):
        parsed = urlparse(self.path)
        length = int(self.headers.get("Content-Length", "0"))
        body = self.rfile.read(length) if length > 0 else b""

        if parsed.path == "/api/reviews":
            try:
                incoming = json.loads(body.decode("utf-8"))
                if not isinstance(incoming, dict):
                    raise ValueError("Expected JSON object of reviews")
                merged = save_reviews(incoming)
                return self._send_json({"ok": True, "savedCount": len(merged), "reviews": merged})
            except Exception as exc:
                return self._send_json({"ok": False, "error": str(exc)}, status=400)

        if parsed.path == "/api/rebuild-from-excel":
            try:
                if body:
                    EXCEL_FILE.write_bytes(body)
                subprocess.run([sys.executable, str(BUILD_SCRIPT)], check=True, cwd=str(ROOT))
                projects_data = json.loads(DATA_JSON.read_text(encoding="utf-8"))
                return self._send_json({"ok": True, "data": projects_data})
            except Exception as exc:
                return self._send_json({"ok": False, "error": str(exc)}, status=500)

        return self._send_json({"ok": False, "error": "Unknown endpoint"}, status=404)


def main():
    parser = argparse.ArgumentParser(description="Start Naue Opportunity Intelligence Portal")
    parser.add_argument("--port", type=int, default=8080, help="Port number (default: 8080)")
    parser.add_argument("--host", type=str, default="127.0.0.1", help="Host address")
    parser.add_argument("--open", action="store_true", help="Open browser automatically")
    args = parser.parse_args()

    mimetypes.add_type("application/javascript", ".js")
    mimetypes.add_type("text/css", ".css")

    url = f"http://{args.host}:{args.port}"
    server = ThreadingHTTPServer((args.host, args.port), TrackerRequestHandler)
    print("=" * 68)
    print("  NAUE GEOSYNTHETICS — PROJECT OPPORTUNITY INTELLIGENCE PORTAL")
    print(f"  Running at: {url}")
    print("  Press Ctrl+C to stop.")
    print("=" * 68)

    if args.open:
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()

    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nShutting down server...")
        server.server_close()


if __name__ == "__main__":
    main()

